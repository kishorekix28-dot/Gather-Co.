from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor

OUTPUT_FILE = r"C:\Users\KISHORE\Desktop\kishore-portfolio\projects\Working_and_Connections_of_3-Axis_MEMS_Gyroscope_and_Mixed_Step-Down_Transformer.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(10, 20, 35)
ACCENT = RGBColor(61, 126, 255)
ACCENT_2 = RGBColor(99, 191, 219)
CARD = RGBColor(21, 35, 52)
TEXT = RGBColor(240, 246, 255)
MUTED = RGBColor(171, 184, 203)
GREEN = RGBColor(61, 180, 120)
YELLOW = RGBColor(246, 190, 99)
RED = RGBColor(223, 100, 100)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide):
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.38))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title(slide, title, subtitle=None):
    set_bg(slide)
    add_top_bar(slide)
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(11.5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.name = 'Aptos'
    run.font.color.rgb = TEXT
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(10.5), Inches(0.5))
        tf2 = sub.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(15)
        r2.font.name = 'Aptos'
        r2.font.color.rgb = MUTED


def add_bullet_box(slide, x, y, w, h, title, bullets):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CARD; box.line.color.rgb = ACCENT; box.line.width = Pt(1.2)
    tb = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.42))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(18); r.font.bold = True; r.font.name = 'Aptos'; r.font.color.rgb = TEXT
    body = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.55), w - Inches(0.5), h - Inches(0.82))
    tf2 = body.text_frame
    tf2.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.level = 0
        p.bullet = True
        p.text = item
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(14); r.font.name = 'Aptos'; r.font.color.rgb = TEXT


def add_card(slide, x, y, w, h, title, body):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CARD; box.line.color.rgb = ACCENT_2; box.line.width = Pt(1.2)
    tb = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.12), w - Inches(0.3), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    for r in p.runs:
        r.font.size = Pt(16); r.font.bold = True; r.font.name = 'Aptos'; r.font.color.rgb = TEXT
    body_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.48), w - Inches(0.3), h - Inches(0.55))
    p2 = body_box.text_frame.paragraphs[0]
    p2.text = body
    for r in p2.runs:
        r.font.size = Pt(11); r.font.name = 'Aptos'; r.font.color.rgb = TEXT

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide)
# title text
header = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(10.5), Inches(1))
p = header.text_frame.paragraphs[0]
p.text = 'Working and Connections of 3-Axis MEMS Gyroscope & Mixed Step-Down Transformer'
for r in p.runs:
    r.font.size = Pt(28); r.font.bold = True; r.font.name = 'Aptos'; r.font.color.rgb = TEXT
sub = slide.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(8.5), Inches(0.5))
p2 = sub.text_frame.paragraphs[0]
p2.text = 'Precision Sensing and Power Regulation in Modern Systems'
for r in p2.runs:
    r.font.size = Pt(18); r.font.name = 'Aptos'; r.font.color.rgb = MUTED
presenter = slide.shapes.add_textbox(Inches(0.75), Inches(6.5), Inches(5), Inches(0.4))
p3 = presenter.text_frame.paragraphs[0]
p3.text = 'Presenter Name: [Your Name]'
for r in p3.runs:
    r.font.size = Pt(14); r.font.name = 'Aptos'; r.font.color.rgb = ACCENT_2
# decorative shapes
circuit = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(9.3), Inches(1.8), Inches(3.0), Inches(3.8))
circuit.fill.solid(); circuit.fill.fore_color.rgb = CARD; circuit.line.color.rgb = ACCENT_2; circuit.line.width = Pt(1.5)
# simple row of nodes
for x in [9.7, 10.7, 11.7]:
    node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(2.4), Inches(0.35), Inches(0.35))
    node.fill.solid(); node.fill.fore_color.rgb = GREEN; node.line.fill.background()
# transformer-like shape
tr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(10.1), Inches(3.3), Inches(1.8), Inches(1.2))
tr.fill.solid(); tr.fill.fore_color.rgb = YELLOW; tr.line.fill.background()
tr_text = slide.shapes.add_textbox(Inches(10.15), Inches(3.66), Inches(1.7), Inches(0.55))
p4 = tr_text.text_frame.paragraphs[0]; p4.text = 'Transformer';
for r in p4.runs:
    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = BG; r.font.name = 'Aptos'

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Overview of Components', 'The sensor and power unit')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(5.7), Inches(3.8), 'The Sensor: 3-Axis MEMS Digital Gyroscope', [
    'Detects angular velocity on X, Y, and Z axes.',
    'Essential for drones, stabilization, and gaming.',
    'Measures rotational motion precisely in real time.'
])
add_bullet_box(slide, Inches(6.9), Inches(2.0), Inches(5.5), Inches(3.8), 'The Power/Actuator Component: Mixed Step-Down Transformer', [
    'Converts high-voltage AC into multiple lower-voltage AC outputs.',
    'Provides the energy source for the control system.',
    'Supports different voltage levels for connected devices.'
])
# mini figs
sensor = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1.4), Inches(5.5), Inches(1.9), Inches(0.9))
sensor.fill.solid(); sensor.fill.fore_color.rgb = ACCENT; sensor.line.fill.background()
text1 = slide.shapes.add_textbox(Inches(1.5), Inches(5.7), Inches(1.7), Inches(0.4)); p5 = text1.text_frame.paragraphs[0]; p5.text='MEMS IC';
for r in p5.runs: r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb=TEXT; r.font.name='Aptos'
tr2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(9.5), Inches(5.5), Inches(2.4), Inches(0.9))
tr2.fill.solid(); tr2.fill.fore_color.rgb = YELLOW; tr2.line.fill.background()
text2 = slide.shapes.add_textbox(Inches(9.8), Inches(5.7), Inches(1.9), Inches(0.4)); p6 = text2.text_frame.paragraphs[0]; p6.text='Multi-tap Transformer';
for r in p6.runs: r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb=BG; r.font.name='Aptos'

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Working of 3-Axis MEMS Gyroscope', 'Core principle is the Coriolis Effect')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(6.0), Inches(4.0), 'Principle', [
    'The Coriolis Effect is the core operating principle.',
    'Tiny silicon structures vibrate at a high frequency inside the chip.',
    'When the device rotates, Coriolis force causes a displacement of these structures.',
    'This movement changes electrical capacitance between the moving and fixed parts.'
])
add_bullet_box(slide, Inches(7.2), Inches(2.0), Inches(5.3), Inches(4.0), 'Output', [
    'The capacitance variation is converted by an onboard ASIC.',
    'Digital data is generated as angular velocity values in degrees/sec.',
    'These values are used to detect rotation on the X, Y, and Z axes.',
    'The result is precise orientation and stabilization feedback.'
])
# small diagram
diagram = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(3.2), Inches(5.0), Inches(2.2), Inches(0.9)); diagram.fill.solid(); diagram.fill.fore_color.rgb=ACCENT; diagram.line.fill.background();
pp = diagram.text_frame.paragraphs[0]; pp.text='Proof Mass';
for r in pp.runs: r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=TEXT; r.font.name='Aptos'
arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(5.5), Inches(5.15), Inches(1.2), Inches(0.5)); arrow.fill.solid(); arrow.fill.fore_color.rgb=GREEN; arrow.line.fill.background()
# Add torque text
note = slide.shapes.add_textbox(Inches(7.0), Inches(5.2), Inches(3.3), Inches(0.5)); p7 = note.text_frame.paragraphs[0]; p7.text='Coriolis Force';
for r in p7.runs: r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=YELLOW; r.font.name='Aptos'

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Connections of the MEMS Gyroscope', 'Most digital gyroscopes use I2C')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(5.6), Inches(3.8), 'Pin Functions', [
    'VCC: Power input (3.3V or 5V DC).',
    'GND: Ground reference.',
    'SCL: Serial clock signal for synchronization.',
    'SDA: Serial data line for sensor output.',
    'AD0: Address select pin for multiple sensors.'
])
# schematic box
schematic = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(7.2), Inches(2.2), Inches(4.8), Inches(3.5))
schematic.fill.solid(); schematic.fill.fore_color.rgb = CARD; schematic.line.color.rgb = ACCENT; schematic.line.width = Pt(1.2)
# draw wires-like lines in a schematic style using simple rectangles
for y in [3.1, 3.9, 4.7]:
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8), Inches(y), Inches(2.0), Inches(0.05))
    line.fill.solid(); line.fill.fore_color.rgb = MUTED; line.line.fill.background()
# microcontroller chip
chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(10.3), Inches(2.9), Inches(1.3), Inches(1.5)); chip.fill.solid(); chip.fill.fore_color.rgb = GREEN; chip.line.fill.background()
chiptxt = slide.shapes.add_textbox(Inches(10.4), Inches(3.5), Inches(1.1), Inches(0.5)); ctp = chiptxt.text_frame.paragraphs[0]; ctp.text='MCU';
for r in ctp.runs: r.font.size=Pt(14); r.font.bold=True; r.font.color.rgb=TEXT; r.font.name='Aptos'
# sensor chip
sensor2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.1), Inches(3.3), Inches(1.1), Inches(1.0)); sensor2.fill.solid(); sensor2.fill.fore_color.rgb = ACCENT_2; sensor2.line.fill.background()
stxt = slide.shapes.add_textbox(Inches(8.15), Inches(3.6), Inches(1.0), Inches(0.4)); sp = stxt.text_frame.paragraphs[0]; sp.text='Gyro';
for r in sp.runs: r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=BG; r.font.name='Aptos'
# label tokens
lab = slide.shapes.add_textbox(Inches(7.8), Inches(5.8), Inches(4.4), Inches(0.5)); p8 = lab.text_frame.paragraphs[0]; p8.text='SCL  SDA  VCC  GND  AD0';
for r in p8.runs: r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb = MUTED; r.font.name='Aptos'

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Working of Mixed Step-Down Transformer', 'Based on mutual induction and Faraday’s Law')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(6.0), Inches(4.2), 'Operating Principle', [
    'Mutual induction is the fundamental concept.',
    'AC current in the primary coil creates a magnetic field in the core.',
    'This induces a voltage in the secondary coil.',
    'A step-down effect occurs because the secondary has fewer turns.',
    'A mixed transformer has multiple taps to provide different output voltages.'
])
# transformer core
core = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.3), Inches(2.4), Inches(3.8), Inches(2.2)); core.fill.solid(); core.fill.fore_color.rgb = CARD; core.line.color.rgb = ACCENT; core.line.width = Pt(1.2)
primary = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.8), Inches(2.8), Inches(0.6), Inches(1.5)); primary.fill.solid(); primary.fill.fore_color.rgb=YELLOW; primary.line.fill.background()
secondary1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(10.1), Inches(2.5), Inches(0.6), Inches(0.7)); secondary1.fill.solid(); secondary1.fill.fore_color.rgb=GREEN; secondary1.line.fill.background()
secondary2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(10.1), Inches(3.4), Inches(0.6), Inches(0.7)); secondary2.fill.solid(); secondary2.fill.fore_color.rgb=GREEN; secondary2.line.fill.background()
secondary3 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(10.1), Inches(4.3), Inches(0.6), Inches(0.7)); secondary3.fill.solid(); secondary3.fill.fore_color.rgb=GREEN; secondary3.line.fill.background()
text3 = slide.shapes.add_textbox(Inches(9.3), Inches(5.1), Inches(2.4), Inches(0.5)); p9 = text3.text_frame.paragraphs[0]; p9.text='6V  12V  24V taps';
for r in p9.runs: r.font.size=Pt(14); r.font.bold=True; r.font.color.rgb=MUTED; r.font.name='Aptos'

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Connections of a Mixed Transformer', 'Primary and secondary wiring')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(5.8), Inches(4.0), 'Primary Side', [
    'Connected to the AC mains (110V/220V).',
    'This side receives the high-voltage input supply.',
    'Proper isolation and insulation are essential for safety.'
])
add_bullet_box(slide, Inches(6.9), Inches(2.0), Inches(5.4), Inches(4.0), 'Secondary Side (Mixed Taps)', [
    '0V: common reference / shared ground.',
    'Tap 1: e.g., 6V for logic boards and sensors.',
    'Tap 2: e.g., 12V for relays or small motors.',
    'Fuse: Always use a safety fuse on the primary side.'
])
# wiring diagram
wiring = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(2.8), Inches(5.5), Inches(7.6), Inches(0.8)); wiring.fill.solid(); wiring.fill.fore_color.rgb = CARD; wiring.line.color.rgb=ACCENT; wiring.line.width=Pt(1.2)
wtxt = slide.shapes.add_textbox(Inches(3.2), Inches(5.72), Inches(6.8), Inches(0.4)); p10 = wtxt.text_frame.paragraphs[0]; p10.text='Primary AC input  →  Transformer  →  0V / 6V / 12V / 24V taps';
for r in p10.runs: r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=TEXT; r.font.name='Aptos'

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Connecting Sensor to Transformer (The Bridge)', 'AC to DC conversion and regulation')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(5.5), Inches(4.2), 'Problem', [
    'The transformer produces AC, but the gyroscope requires DC power.',
    'A direct connection would not work for the sensor electronics.'
])
add_bullet_box(slide, Inches(6.6), Inches(2.0), Inches(5.8), Inches(4.2), 'Connection Path', [
    'Transformer Tap: 6V AC output.',
    'Bridge Rectifier: Converts AC into pulsating DC.',
    'Filter Capacitor: Smooths the pulsating supply.',
    'Voltage Regulator: Provides a steady 5V DC to the gyroscope.'
])
# flow chart
flow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(3.5), Inches(5.3), Inches(6.0), Inches(0.7)); flow.fill.solid(); flow.fill.fore_color.rgb = CARD; flow.line.color.rgb = ACCENT; flow.line.width = Pt(1.2)
ft = slide.shapes.add_textbox(Inches(3.7), Inches(5.52), Inches(5.7), Inches(0.35)); pf = ft.text_frame.paragraphs[0]; pf.text='Transformer → Rectifier → Filter Capacitor → Voltage Regulator → MEMS Gyroscope';
for r in pf.runs: r.font.size=Pt(14); r.font.bold=True; r.font.color.rgb=TEXT; r.font.name='Aptos'

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Practical Application Example', 'Camera gimbal stabilizer system')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(5.5), Inches(3.8), 'System Role of Transformer', [
    'Powers motors and the control board.',
    'Provides stable source voltages for multiple functions.'
])
add_bullet_box(slide, Inches(6.7), Inches(2.0), Inches(5.6), Inches(3.8), 'System Role of Gyroscope', [
    'Detects camera tilting and angular motion.',
    'Sends orientation feedback to the controller.',
    'The controller drives motors to resist unwanted motion.'
])
# gimbal-like shape
gimbal = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(4.8), Inches(5.5), Inches(1.4), Inches(1.4)); gimbal.fill.solid(); gimbal.fill.fore_color.rgb=ACCENT_2; gimbal.line.fill.background()
# simple ring diagram
ring = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(8.9), Inches(5.5), Inches(1.8), Inches(1.8)); ring.fill.solid(); ring.fill.fore_color.rgb=YELLOW; ring.line.fill.background()

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Precautions and Troubleshooting', 'Reliable operation and safety')
add_bullet_box(slide, Inches(0.8), Inches(2.0), Inches(5.8), Inches(3.8), 'MEMS Gyroscope Precautions', [
    'Sensitive to mechanical vibrations; mount on a damping pad.',
    'Keep sensor cabling short and shielded if possible.',
    'Avoid noisy power lines near the signal pins.'
])
add_bullet_box(slide, Inches(6.9), Inches(2.0), Inches(5.5), Inches(3.8), 'Transformer Precautions', [
    'Use proper insulation and heat dissipation.',
    'Keep it away from digital sensor wires to reduce EMI.',
    'Ensure airflow and avoid overheating.',
    'Install a fuse on the primary side for safety.'
])
# icon row maybe
a = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(2.2), Inches(5.9), Inches(0.7), Inches(0.7)); a.fill.solid(); a.fill.fore_color.rgb = GREEN; a.line.fill.background()
b = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(8.7), Inches(5.9), Inches(0.7), Inches(0.7)); b.fill.solid(); b.fill.fore_color.rgb=RED; b.line.fill.background()

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide); add_top_bar(slide); add_title(slide, 'Conclusion', 'Final summary')
add_bullet_box(slide, Inches(1.1), Inches(2.0), Inches(5.5), Inches(4.0), 'System Roles', [
    'The 3-Axis MEMS Gyroscope is the “eyes” of the system, sensing orientation.',
    'The Mixed Step-Down Transformer is the “heart,” providing varied power levels.'
])
add_bullet_box(slide, Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.0), 'Key Insight', [
    'Integrating them correctly requires clean AC-to-DC conversion and regulated voltages.',
    'This ensures both sensing accuracy and stable power delivery.',
    'The combined design enables reliable control in modern systems.'
])
qs = slide.shapes.add_textbox(Inches(5.7), Inches(6.0), Inches(2.0), Inches(0.7)); qsp = qs.text_frame.paragraphs[0]; qsp.text='Questions?';
for r in qsp.runs: r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=ACCENT_2; r.font.name='Aptos'

prs.save(OUTPUT_FILE)
print(f"Presentation created successfully: {OUTPUT_FILE}")
